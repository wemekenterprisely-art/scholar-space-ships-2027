"""excel_generator - 5-sheet scholarship workbook (CareerOps quality)."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

# Sheet 1: All Scholarships
ALL_COLS = ["#", "Title", "University", "Country", "Level", "Funding", "Deadline",
            "English Req", "Score", "Source", "Link", "Notes"]

# Sheet 2: Fresh Matches (accumulated)
FRESH_COLS = ALL_COLS + ["AI Verdict", "First Seen", "Applied"]

# Sheet 3: Applications
APP_COLS = ["Title", "Link", "University", "Country", "Status", "Deadline", "Notes", "Updated"]

# Sheet 4: Deadlines & Notes
DEADLINE_COLS = ["Title", "Link", "Deadline", "Days Left", "Action", "Funding", "Notes"]

# Sheet 5: Daily Log
LOG_COLS = ["Date", "Time", "Fetched", "Candidates", "Matches", "New", "Seconds", "Sources"]


def _esc(s):
    return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def extract_uni(s):
    import re
    university = s.get("university", s.get("institution", ""))
    if not university:
        title = s.get("title", "")
        patterns = [
            r"at\s+([\w\s,]+(?:University|Institute|College|School)[\w\s]*)",
            r"([\w\s,]+(?:University|Institute|College|School))",
            r"(CSC|Chevening|Fulbright|DAAD|Erasmus|MEXT|Turkiye|Stipendium)\s+Scholarship",
        ]
        for p in patterns:
            m = re.search(p, title, re.I)
            if m:
                university = m.group(1).strip() if m.lastindex else m.group(0).strip()
                university = re.sub(r"\s*\d{4}\s*$", "", university).strip()
                break
    return university


def get_recommendation(score):
    if score >= 85: return "STRONG MATCH - Clear fit. Apply now."
    if score >= 75: return "GOOD MATCH - Strong overlap. Review and apply."
    if score >= 65: return "MODERATE MATCH - Review requirements before applying."
    return "LOW MATCH - May not fit profile."


def _write_header(ws, cols, widths):
    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center"), THIN_BORDER
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def _write_rows(ws, rows, start=2):
    for i, r in enumerate(rows, start=start):
        for j, val in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.alignment, c.border = WRAP, THIN_BORDER
        ws.row_dimensions[i].height = 45


def make_workbook(candidates, matches, app_track, scan_hist):
    wb = Workbook()

    # ---- Sheet 1: All Scholarships ----
    ws = wb.active
    ws.title = "All Scholarships"
    _write_header(ws, ALL_COLS, [5, 50, 30, 14, 10, 14, 14, 12, 8, 14, 45, 40])

    rows = []
    for i, s in enumerate(sorted(candidates, key=lambda x: x.get("final_score", 0), reverse=True), start=1):
        uni = extract_uni(s)
        notes = " | ".join(s.get("notes", []))
        ielts = s.get("english_requirement", "unknown")
        ielts_label = "Not Required" if ielts == "exempt" else ("Required" if ielts == "required" else "Verify")
        rows.append([
            i, s.get("title", ""), uni or "Check listing", s.get("country", "Open/Global"),
            s.get("level", "Unknown"), s.get("funding", "Unknown"), s.get("deadline", "N/A"),
            ielts_label, s.get("final_score", s.get("det_score", 0)),
            s.get("source", ""), s.get("url", ""), notes
        ])
    _write_rows(ws, rows)
    # Hyperlinks
    for r in range(2, len(rows) + 2):
        c = ws.cell(row=r, column=11)
        if c.value:
            c.hyperlink, c.font = c.value, LINK_FONT

    # ---- Sheet 2: Fresh Matches (accumulated) ----
    ws2 = wb.create_sheet("Fresh Matches")
    _write_header(ws2, FRESH_COLS, [5, 50, 30, 14, 10, 14, 14, 12, 8, 14, 45, 40, 30, 14, 10])

    rows2 = []
    for i, s in enumerate(sorted(matches, key=lambda x: x.get("final_score", 0), reverse=True), start=1):
        uni = extract_uni(s)
        notes = " | ".join(s.get("notes", []))
        ielts = s.get("english_requirement", "unknown")
        ielts_label = "Not Required" if ielts == "exempt" else ("Required" if ielts == "required" else "Verify")
        applied = "Yes" if s.get("applied") else "No"
        rows2.append([
            i, s.get("title", ""), uni or "Check listing", s.get("country", "Open/Global"),
            s.get("level", "Unknown"), s.get("funding", "Unknown"), s.get("deadline", "N/A"),
            ielts_label, s.get("final_score", s.get("det_score", 0)),
            s.get("source", ""), s.get("url", ""), notes,
            s.get("ai_verdict", ""), s.get("first_seen", "")[:10], applied
        ])
    _write_rows(ws2, rows2)
    for r in range(2, len(rows2) + 2):
        c = ws2.cell(row=r, column=11)
        if c.value:
            c.hyperlink, c.font = c.value, LINK_FONT

    # ---- Sheet 3: Applications ----
    ws3 = wb.create_sheet("Applications")
    _write_header(ws3, APP_COLS, [50, 45, 30, 14, 14, 14, 40, 20])

    rows3 = []
    for sid, a in (app_track or {}).items():
        rows3.append([a.get("title", sid), a.get("url", ""), a.get("university", ""),
                      a.get("country", ""), a.get("status", "Apply"), a.get("deadline", ""),
                      a.get("notes", ""), a.get("updated", "")])
    _write_rows(ws3, rows3)
    for r in range(2, len(rows3) + 2):
        c = ws3.cell(row=r, column=2)
        if c.value:
            c.hyperlink, c.font = c.value, LINK_FONT

    # ---- Sheet 4: Deadlines & Notes ----
    ws4 = wb.create_sheet("Deadlines & Notes")
    _write_header(ws4, DEADLINE_COLS, [50, 45, 14, 10, 14, 14, 40])

    rows4 = []
    for s in sorted(matches, key=lambda x: x.get("deadline_days") if x.get("deadline_days") is not None else 9999):
        d = s.get("deadline_days")
        action = "APPLY NOW" if (d is not None and d <= 21) else ("This month" if d is not None and d <= 45 else "Plan ahead")
        rows4.append([s.get("title", ""), s.get("url", ""), s.get("deadline", "N/A"),
                      d if d is not None else "N/A", action, s.get("funding", "Unknown"),
                      " | ".join(s.get("notes", []))])
    _write_rows(ws4, rows4)
    for r in range(2, len(rows4) + 2):
        c = ws4.cell(row=r, column=2)
        if c.value:
            c.hyperlink, c.font = c.value, LINK_FONT

    # ---- Sheet 5: Daily Log ----
    ws5 = wb.create_sheet("Daily Log")
    _write_header(ws5, LOG_COLS, [12, 24, 10, 12, 10, 8, 10, 60])

    rows5 = [[h.get("date"), h.get("time"), h.get("fetched"), h.get("candidates"),
              h.get("matches"), h.get("new"), h.get("seconds"),
              json.dumps(h.get("sources", {}), ensure_ascii=False)] for h in scan_hist[-30:]]
    _write_rows(ws5, rows5)

    return wb
